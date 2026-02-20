# -*- coding: utf-8 -*-
import json
import os
import random
import re
import threading
import time
import unicodedata
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import Workbook
from selenium.common.exceptions import WebDriverException
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

try:
    from app.services.authorization import authorize_action, build_auth_audit_event
    from app.services.audit_store import append_audit_event
    from app.services.download_contract import execute_with_retry_contract
    from app.services.runtime_policy import load_runtime_policy_gate
    from app.services.search_stability import (
        build_search_context,
        stable_sort_tenders,
        validate_download_scope,
    )
    from app.services.template_builder import extract_placeholders_from_docx, render_docx_template
    from app.services.ux_guidance import build_corrective_guidance
    from app.services.validation_engine import validate_required_inputs
    from app.services.workflow_router import route_action
    from app.services.workspace_pack import create_workspace_pack
    from app.services.tender_search import (
        TenderRow,
        click_download_all_in_modal,
        collect_tenders,
        collect_all_pages,
        ensure_on_notices,
        find_dossier_on_pages,
        has_any_result_rows,
        handle_download_doc_without_login,
        login_on_download_doc,
        open_search_panel,
        search_keyword,
        setup_driver,
        wait_for_result_rows,
    )
except ImportError:
    from services.authorization import authorize_action, build_auth_audit_event
    from services.audit_store import append_audit_event
    from services.download_contract import execute_with_retry_contract
    from services.runtime_policy import load_runtime_policy_gate
    from services.search_stability import (
        build_search_context,
        stable_sort_tenders,
        validate_download_scope,
    )
    from services.template_builder import extract_placeholders_from_docx, render_docx_template
    from services.ux_guidance import build_corrective_guidance
    from services.validation_engine import validate_required_inputs
    from services.workflow_router import route_action
    from services.workspace_pack import create_workspace_pack
    from services.tender_search import (
        TenderRow,
        click_download_all_in_modal,
        collect_tenders,
        collect_all_pages,
        ensure_on_notices,
        find_dossier_on_pages,
        has_any_result_rows,
        handle_download_doc_without_login,
        login_on_download_doc,
        open_search_panel,
        search_keyword,
        setup_driver,
        wait_for_result_rows,
    )

MATRIX_BG = "#030A03"
MATRIX_PANEL = "#071207"
MATRIX_PANEL_ALT = "#0A170A"
MATRIX_FG = "#8CFF8C"
MATRIX_ACCENT = "#00FF41"
MATRIX_MUTED = "#5ECB74"
COMPLIANCE_RULES_DIR = Path.cwd() / "compliance" / "rules"


class TenderSearchFrame(ttk.Frame):
    def __init__(self, master):
        super().__init__(master)
        self.var_keyword = tk.StringVar(value="Internet")
        self.var_download = tk.StringVar(value=str(Path.cwd() / "downloads"))
        self.var_headless = tk.BooleanVar(value=True)
        self.var_username = tk.StringVar(value="")
        self.var_password = tk.StringVar(value="")
        self.var_role = tk.StringVar(value="procurement_officer")
        self.var_process_mode = tk.StringVar(value="esjn")
        self.var_search_mode = tk.StringVar(value="Mode: idle")
        self.var_search_quality = tk.StringVar(value="Quality: raw=0 filtered=0 fallback=no pages=0")
        # Default to first page for faster interactive feedback; users can enable full pagination.
        self.var_collect_all_pages = tk.BooleanVar(value=False)
        self.var_max_pages = tk.StringVar(value="5")
        self.var_strict_filter = tk.BooleanVar(value=True)
        self.var_match_mode = tk.StringVar(value="contains")
        self.results: list[TenderRow] = []
        self.last_search_context: dict | None = None
        self.audit_file = Path.cwd() / "compliance" / "audit" / "events.jsonl"
        self.driver = None
        self.wait = None
        self._driver_lock = threading.Lock()
        self._build_ui()

    def _build_ui(self):
        top = ttk.Frame(self)
        top.pack(fill="x", padx=8, pady=8)

        ttk.Label(top, text="Keyword:").grid(row=0, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.var_keyword, width=24).grid(row=0, column=1, sticky="w")
        ttk.Label(top, text="Username (optional):").grid(row=0, column=2, sticky="w", padx=(10, 0))
        ttk.Entry(top, textvariable=self.var_username, width=20).grid(row=0, column=3, sticky="w")
        ttk.Label(top, text="Password (optional):").grid(row=0, column=4, sticky="w", padx=(10, 0))
        ttk.Entry(top, textvariable=self.var_password, show="*", width=20).grid(row=0, column=5, sticky="w")
        ttk.Label(top, text="Role:").grid(row=0, column=6, sticky="w", padx=(10, 0))
        role_box = ttk.Combobox(
            top,
            textvariable=self.var_role,
            values=("procurement_officer", "admin", "compliance_auditor", "viewer"),
            width=20,
            state="readonly",
        )
        role_box.grid(row=0, column=7, sticky="w")
        ttk.Label(top, text="Process mode:").grid(row=0, column=8, sticky="w", padx=(10, 0))
        mode_box = ttk.Combobox(
            top,
            textvariable=self.var_process_mode,
            values=("esjn", "epazar"),
            width=12,
            state="readonly",
        )
        mode_box.grid(row=0, column=9, sticky="w")

        ttk.Label(top, text="Download folder:").grid(row=1, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.var_download, width=60).grid(row=1, column=1, columnspan=4, sticky="we")
        ttk.Button(top, text="Browse...", command=self.choose_dir).grid(row=1, column=5, sticky="w")
        ttk.Checkbutton(top, text="Headless (always on)", variable=self.var_headless, state="disabled").grid(
            row=1, column=6, sticky="w", padx=(8, 0)
        )
        ttk.Checkbutton(top, text="Collect all pages", variable=self.var_collect_all_pages).grid(
            row=2, column=0, sticky="w"
        )
        ttk.Label(top, text="Max pages:").grid(row=2, column=1, sticky="e")
        ttk.Entry(top, textvariable=self.var_max_pages, width=6).grid(row=2, column=2, sticky="w")
        ttk.Checkbutton(top, text="Strict keyword filter", variable=self.var_strict_filter).grid(
            row=2, column=3, sticky="w", padx=(10, 0)
        )
        ttk.Label(top, text="Match mode:").grid(row=2, column=4, sticky="e")
        mode_box = ttk.Combobox(
            top,
            textvariable=self.var_match_mode,
            values=("contains", "all_words", "exact_phrase", "regex"),
            width=14,
            state="readonly",
        )
        mode_box.grid(row=2, column=5, sticky="w")

        btns = ttk.Frame(self)
        btns.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Button(btns, text="Connect", command=self.on_connect).pack(side="left")
        ttk.Button(btns, text="Search", command=self.on_search).pack(side="left")
        ttk.Button(btns, text="Download selected", command=self.on_download_selected).pack(
            side="left", padx=6
        )
        ttk.Button(btns, text="Export Excel", command=self.on_export_excel).pack(side="left")
        ttk.Button(btns, text="Copy Logs", command=self.copy_logs).pack(side="left", padx=6)
        ttk.Label(btns, textvariable=self.var_search_mode).pack(side="left", padx=(12, 0))
        ttk.Label(btns, textvariable=self.var_search_quality).pack(side="left", padx=(12, 0))

        cols = ("#", "Title", "Institution", "Deadline", "DossierID")
        wrap = ttk.Frame(self)
        wrap.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        self.tree = ttk.Treeview(wrap, columns=cols, show="headings", selectmode="extended")
        widths = (50, 420, 260, 140, 260)
        for col, width in zip(cols, widths):
            self.tree.heading(col, text=col)
            self.tree.column(col, width=width, anchor="w")
        self.tree.pack(side="left", fill="both", expand=True)
        yscroll = ttk.Scrollbar(wrap, orient="vertical", command=self.tree.yview)
        yscroll.pack(side="right", fill="y")
        self.tree.configure(yscroll=yscroll.set)

        self.log_text = tk.Text(self, height=10, wrap="word")
        self.log_text.pack(fill="x", padx=8, pady=(0, 8))
        self.log_text.configure(
            bg=MATRIX_BG,
            fg=MATRIX_FG,
            insertbackground=MATRIX_ACCENT,
            selectbackground=MATRIX_ACCENT,
            selectforeground=MATRIX_BG,
            highlightthickness=1,
            highlightbackground=MATRIX_ACCENT,
            relief="flat",
        )
        self.log_text.config(state="disabled")
        log_actions = ttk.Frame(self)
        log_actions.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Button(log_actions, text="Copy Logs", command=self.copy_logs).pack(side="left")

    def log(self, msg: str):
        self.log_text.config(state="normal")
        self.log_text.insert("end", time.strftime("%H:%M:%S  ") + msg + "\n")
        self.log_text.see("end")
        self.log_text.config(state="disabled")

    def copy_logs(self):
        text = self.log_text.get("1.0", "end").strip()
        if not text:
            messagebox.showinfo("Logs", "No logs to copy.")
            return
        self.clipboard_clear()
        self.clipboard_append(text)
        self.log("INFO: Logs copied to clipboard.")

    def _enforce_runtime_policy(self, action: str) -> bool:
        decision = load_runtime_policy_gate(COMPLIANCE_RULES_DIR).decide(action)
        active = ",".join(decision.active_rule_ids) if decision.active_rule_ids else "none"
        self.log(
            f"POLICY_GATE action={action} module={decision.module} "
            f"allowed={str(decision.allowed).lower()} active_rules={active} reason={decision.reason}"
        )
        if not decision.allowed:
            messagebox.showerror("Compliance gate", decision.reason)
            return False
        return True

    def choose_dir(self):
        selected = filedialog.askdirectory(initialdir=self.var_download.get() or str(Path.cwd()))
        if selected:
            self.var_download.set(selected)

    def ensure_driver(self):
        with self._driver_lock:
            if self.driver is None:
                # Headless mode is enforced for every browser action.
                self.var_headless.set(True)
                self.driver = setup_driver(True, self.var_download.get().strip())
                self.wait = WebDriverWait(self.driver, 20)
        return self.driver, self.wait

    def on_connect(self):
        def work():
            t0 = time.perf_counter()
            try:
                driver, wait = self.ensure_driver()
                ensure_on_notices(driver, wait)
                self.log(f"INFO: Connected. ready_sec={time.perf_counter() - t0:.2f}")
            except Exception as exc:
                self.log(f"ERROR: Connect failed: {exc}")

        threading.Thread(target=work, daemon=True).start()

    @staticmethod
    def _normalize_token_text(value: str) -> str:
        text = unicodedata.normalize("NFKC", value or "").lower()
        text = re.sub(r"[^0-9a-zа-ш]+", " ", text, flags=re.IGNORECASE)
        return re.sub(r"\s+", " ", text).strip()

    @staticmethod
    def _cyr_to_lat(value: str) -> str:
        table = str.maketrans(
            {
                "а": "a",
                "б": "b",
                "в": "v",
                "г": "g",
                "д": "d",
                "ѓ": "gj",
                "е": "e",
                "ж": "zh",
                "з": "z",
                "ѕ": "dz",
                "и": "i",
                "ј": "j",
                "к": "k",
                "л": "l",
                "љ": "lj",
                "м": "m",
                "н": "n",
                "њ": "nj",
                "о": "o",
                "п": "p",
                "р": "r",
                "с": "s",
                "т": "t",
                "ќ": "kj",
                "у": "u",
                "ф": "f",
                "х": "h",
                "ц": "c",
                "ч": "ch",
                "џ": "dj",
                "ш": "sh",
            }
        )
        return (value or "").translate(table)

    def _keyword_match(self, text: str, keyword: str) -> bool:
        kw = (keyword or "").strip()
        if not kw:
            return True
        hay = (text or "")
        hay_norm = self._normalize_token_text(hay)
        hay_lat = self._normalize_token_text(self._cyr_to_lat(hay_norm))
        mode = (self.var_match_mode.get() or "contains").strip()
        kw_variants = {kw.lower()}
        if kw.lower() == "internet":
            kw_variants.add("интернет")
        if kw.lower() == "интернет":
            kw_variants.add("internet")
        kw_norm_variants = {self._normalize_token_text(v) for v in kw_variants if v.strip()}
        kw_lat_variants = {self._normalize_token_text(self._cyr_to_lat(v)) for v in kw_norm_variants}
        all_variants = {v for v in (kw_norm_variants | kw_lat_variants) if v}
        if mode == "exact_phrase":
            return any(v in hay_norm or v in hay_lat for v in all_variants)
        if mode == "all_words":
            for variant in all_variants:
                words = [w for w in variant.split() if w]
                if all((w in hay_norm) or (w in hay_lat) for w in words):
                    return True
            return False
        if mode == "regex":
            try:
                return re.search(kw, text or "", flags=re.IGNORECASE) is not None
            except re.error:
                self.log("WARN: Invalid regex. Falling back to contains.")
                return any(v in hay_norm or v in hay_lat for v in all_variants)
        return any(v in hay_norm or v in hay_lat for v in all_variants)

    def _post_filter_by_keyword(self, rows: list[TenderRow], keyword: str) -> list[TenderRow]:
        kw = (keyword or "").strip().lower()
        if not kw:
            return rows
        out: list[TenderRow] = []
        for r in rows:
            hay = " | ".join(
                [
                    (r.title or ""),
                    (r.institution or ""),
                    (r.deadline or ""),
                    (r.row_text or ""),
                ]
            )
            if self._keyword_match(hay, keyword):
                out.append(r)
        return out

    def on_search(self):
        keyword = (self.var_keyword.get() or "").strip()
        if not keyword:
            messagebox.showwarning("Missing keyword", "Enter a keyword.")
            return
        if not self._enforce_runtime_policy("search"):
            return
        route = route_action(self.var_process_mode.get(), "search")
        self.log(f"WORKFLOW_ROUTE mode={route.mode} action=search allowed={str(route.allowed).lower()} message={route.message}")
        if not route.allowed:
            messagebox.showerror("Workflow routing", route.message)
            return

        def work():
            self.log(f"SEARCH: {keyword}")
            try:
                used_fallback = False
                try:
                    max_pages = max(1, int((self.var_max_pages.get() or "5").strip()))
                except ValueError:
                    max_pages = 5
                    self.log("WARN: Invalid max pages value. Using 5.")
                driver, wait = self.ensure_driver()
                snapshot_dir = str(Path(self.var_download.get().strip() or str(Path.cwd() / "downloads")) / "debug")
                ensure_on_notices(driver, wait)
                open_search_panel(driver, wait, self.log)
                search_keyword(driver, wait, keyword)
                if self.var_collect_all_pages.get():
                    self.results = collect_all_pages(
                        driver,
                        wait,
                        self.log,
                        max_pages=max_pages,
                        snapshot_dir=snapshot_dir,
                    )
                else:
                    wait_for_result_rows(
                        driver,
                        wait,
                        self.log,
                        attempts=3,
                        base_delay_sec=1.2,
                        snapshot_dir=snapshot_dir,
                        max_total_wait_sec=8.0,
                    )
                    self.results = collect_tenders(driver, wait, self.log)

                if len(self.results) == 0:
                    self.log("INFO: No results after first keyword filter attempt. Retrying once.")
                    # Reset form/page state and retry once.
                    ensure_on_notices(driver, wait)
                    open_search_panel(driver, wait, self.log)
                    search_keyword(driver, wait, keyword)
                    if self.var_collect_all_pages.get():
                        self.results = collect_all_pages(
                            driver,
                            wait,
                            self.log,
                            max_pages=max_pages,
                            snapshot_dir=snapshot_dir,
                        )
                    else:
                        wait_for_result_rows(
                            driver,
                            wait,
                            self.log,
                            attempts=3,
                            base_delay_sec=1.2,
                            snapshot_dir=snapshot_dir,
                            max_total_wait_sec=8.0,
                        )
                        self.results = collect_tenders(driver, wait, self.log)

                if len(self.results) == 0:
                    # Final fallback: collect baseline rows without keyword filter.
                    self.log(
                        "WARN: search_filter_failed=true; collecting baseline results without filter."
                    )
                    ensure_on_notices(driver, wait)
                    if self.var_collect_all_pages.get():
                        self.results = collect_all_pages(
                            driver,
                            wait,
                            self.log,
                            max_pages=max_pages,
                            snapshot_dir=snapshot_dir,
                        )
                    else:
                        self.results = collect_tenders(driver, wait, self.log)
                    used_fallback = True

                raw_count = len(self.results)
                raw_results = list(self.results)
                if self.var_strict_filter.get():
                    self.results = self._post_filter_by_keyword(self.results, keyword)
                filtered_count = len(self.results)
                if self.var_strict_filter.get() and raw_count > 0 and filtered_count == 0:
                    self.log(
                        "INFO: Strict post-filter matched 0 rows. Disable strict filter to inspect raw rows."
                    )
                    for i, sample in enumerate(raw_results[:3], start=1):
                        txt = (sample.row_text or sample.title or "")[:220]
                        self.log(f"DEBUG_FILTER_SAMPLE[{i}]: {txt}")
                self.results = stable_sort_tenders(self.results)
                if filtered_count != raw_count:
                    self.log(
                        f"INFO: Applied strict keyword post-filter: {raw_count} -> {filtered_count} rows."
                    )
                elif not self.var_strict_filter.get():
                    self.log("INFO: Strict keyword filter is OFF; showing unfiltered collected rows.")

                pages_scanned = max((r.source_page for r in self.results), default=0)
                if pages_scanned == 0 and raw_count > 0:
                    pages_scanned = 1

                if used_fallback:
                    self.var_search_mode.set("Mode: fallback baseline (search_filter_failed=true)")
                else:
                    self.var_search_mode.set("Mode: filtered")
                self.var_search_quality.set(
                    f"Quality: raw={raw_count} filtered={filtered_count} "
                    f"fallback={'yes' if used_fallback else 'no'} pages={pages_scanned}"
                )
                self.last_search_context = build_search_context(
                    keyword=keyword,
                    match_mode=self.var_match_mode.get(),
                    strict_filter=self.var_strict_filter.get(),
                    rows=self.results,
                )

                self.tree.delete(*self.tree.get_children())
                for row in self.results:
                    self.tree.insert(
                        "",
                        "end",
                        values=(row.index, row.title, row.institution, row.deadline, row.dossier_id),
                    )
            except WebDriverException as exc:
                self.var_search_mode.set("Mode: error")
                self.log(f"ERROR: WebDriver: {exc}")
            except Exception as exc:
                self.var_search_mode.set("Mode: error")
                self.log(f"ERROR: {exc}")

        threading.Thread(target=work, daemon=True).start()

    def on_download_selected(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showinfo("No selection", "Select one or more tenders.")
            return
        if not self._enforce_runtime_policy("download_selected"):
            return
        selected_dossiers = [str(self.tree.item(i)["values"][4]) for i in selected]
        scope_ok, scope_msg = validate_download_scope(
            context=self.last_search_context,
            current_keyword=(self.var_keyword.get() or "").strip(),
            selected_dossier_ids=selected_dossiers,
        )
        if not scope_ok:
            self.log(f"DOWNLOAD_SCOPE_ERROR: {scope_msg}")
            messagebox.showerror("Download scope", scope_msg)
            return

        username = (self.var_username.get() or "").strip()
        password = self.var_password.get() or ""
        role = (self.var_role.get() or "").strip()
        route = route_action(self.var_process_mode.get(), "download_selected")
        self.log(
            f"WORKFLOW_ROUTE mode={route.mode} action=download_selected "
            f"allowed={str(route.allowed).lower()} message={route.message}"
        )
        if not route.allowed:
            messagebox.showerror("Workflow routing", route.message)
            return
        decision = authorize_action("download_selected", username, role)
        self.log(build_auth_audit_event("download_selected", username, role, decision.allowed, decision.reason))
        if not decision.allowed:
            messagebox.showerror("Authorization denied", decision.reason)
            return

        def work():
            total_started = 0
            try:
                driver, wait = self.ensure_driver()
                keyword = (self.last_search_context or {}).get("keyword", "").strip()
                try:
                    max_pages = max(1, int((self.var_max_pages.get() or "5").strip()))
                except ValueError:
                    max_pages = 5

                def prepare_download_scope() -> bool:
                    # If search results are already present, reuse current context.
                    try:
                        if has_any_result_rows(driver):
                            self.log("INFO: Reusing current results context for download scope.")
                            return True
                    except Exception:
                        pass

                    for attempt in range(1, 4):
                        try:
                            ensure_on_notices(driver, wait)
                            open_search_panel(driver, wait, self.log)
                            if keyword:
                                search_keyword(driver, wait, keyword)
                                found_rows = wait_for_result_rows(
                                    driver,
                                    wait,
                                    self.log,
                                    attempts=3,
                                    base_delay_sec=1.0,
                                    snapshot_dir=str(
                                        Path(
                                            self.var_download.get().strip()
                                            or str(Path.cwd() / "downloads")
                                        )
                                        / "debug"
                                    ),
                                )
                                if not found_rows:
                                    raise RuntimeError("Search scope prepared but no rows became visible.")
                            if not has_any_result_rows(driver):
                                raise RuntimeError("Search scope contains no visible result rows.")
                            return True
                        except Exception as exc:
                            self.log(
                                f"WARN: prepare_download_scope attempt {attempt}/3 failed: "
                                f"{type(exc).__name__}: {exc}"
                            )
                            time.sleep(1.0 * attempt)
                    return False

                for idx, item_id in enumerate(selected, 1):
                    _, title, institution, deadline, dossier_id = self.tree.item(item_id)["values"]
                    visible_dossiers = {
                        str(self.tree.item(i)["values"][4]) for i in self.tree.get_children()
                    }
                    if str(dossier_id) not in visible_dossiers:
                        raise RuntimeError(
                            "Download guard blocked dossier outside current visible filtered rows."
                        )
                    self.log(f"DOWNLOAD [{idx}/{len(selected)}] {title} ({institution}) [{deadline}]")

                    def op(_: int) -> int:
                        found = find_dossier_on_pages(
                            driver, wait, dossier_id, self.log, max_pages=max_pages
                        )
                        if not found:
                            self.log(
                                "INFO: Dossier not found in current context, rebuilding keyword scope."
                            )
                            if not prepare_download_scope():
                                raise RuntimeError("Could not prepare search scope for download.")
                            found = find_dossier_on_pages(
                                driver, wait, dossier_id, self.log, max_pages=max_pages
                            )
                        if not found:
                            raise RuntimeError(
                                f"Dossier not found on first {max_pages} pages in current search scope: {dossier_id}"
                            )
                        click_download_all_in_modal(driver, wait)
                        time.sleep(1.0)

                        if len(driver.window_handles) > 1:
                            driver.switch_to.window(driver.window_handles[-1])

                        started_local = handle_download_doc_without_login(driver, self.log)
                        if started_local == 0 and username and password:
                            if login_on_download_doc(driver, username, password, self.log):
                                try:
                                    all_btn = WebDriverWait(driver, 5).until(
                                        EC.element_to_be_clickable((By.ID, "ctl00_publicAccess_btnDownloadAll"))
                                    )
                                    all_btn.click()
                                    time.sleep(1.2)
                                except Exception:
                                    pass
                                started_local = handle_download_doc_without_login(driver, self.log)
                        if started_local == 0:
                            raise RuntimeError("No direct download links found.")
                        return started_local

                    result = execute_with_retry_contract(
                        operation=op,
                        max_attempts=2,
                        on_event=lambda m, d=dossier_id: self.log(f"DOWNLOAD_STATE dossier={d} {m}"),
                    )
                    if result.status != "success":
                        err_msg = result.error.user_message if result.error else "Unknown download error."
                        guidance = build_corrective_guidance(
                            error_code=(result.error.code if result.error else "unexpected_error"),
                            action="download_selected",
                            mode=self.var_process_mode.get(),
                        )
                        self.log(
                            f"GUIDANCE code={guidance['error_code']} retry_safe={str(guidance['retry_safe']).lower()}"
                        )
                        for step in guidance["steps"]:
                            self.log(f"GUIDANCE_STEP: {step}")
                        append_audit_event(
                            audit_file=self.audit_file,
                            event_type="download_selected",
                            actor=username or "anonymous",
                            module="download",
                            status="failed",
                            dossier_id=str(dossier_id),
                            metadata={
                                "institution": str(institution),
                                "deadline": str(deadline),
                                "attempts_used": result.attempts_used,
                                "error_code": result.error.code if result.error else "unknown",
                                "error_message": err_msg,
                            },
                        )
                        raise RuntimeError(
                            f"Download failed for dossier {dossier_id}: {err_msg} "
                            f"(attempts={result.attempts_used})"
                        )

                    total_started += result.started_count
                    append_audit_event(
                        audit_file=self.audit_file,
                        event_type="download_selected",
                        actor=username or "anonymous",
                        module="download",
                        status="success",
                        dossier_id=str(dossier_id),
                        metadata={
                            "institution": str(institution),
                            "deadline": str(deadline),
                            "attempts_used": result.attempts_used,
                            "started_count": result.started_count,
                        },
                    )
                self.log(f"DONE: Started downloads: {total_started}")
            except Exception as exc:
                self.log(f"ERROR: Download failed: {exc}")
                guidance = build_corrective_guidance(
                    error_code="unexpected_error",
                    action="download_selected",
                    mode=self.var_process_mode.get(),
                )
                self.log(
                    f"GUIDANCE code={guidance['error_code']} retry_safe={str(guidance['retry_safe']).lower()}"
                )
                for step in guidance["steps"]:
                    self.log(f"GUIDANCE_STEP: {step}")

        threading.Thread(target=work, daemon=True).start()

    def on_export_excel(self):
        if not self.results:
            messagebox.showinfo("No data", "Search first.")
            return
        output = filedialog.asksaveasfilename(
            title="Save results",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile="tender-results.xlsx",
        )
        if not output:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Tender Results"
        ws.append(["Row", "Title", "Institution", "Deadline", "DossierID"])
        for row in self.results:
            ws.append([row.index, row.title, row.institution, row.deadline, row.dossier_id])
        wb.save(output)
        self.log(f"EXPORTED: {output}")
        if os.name == "nt":
            os.startfile(output)

    def get_profile_data(self) -> dict:
        return {
            "keyword": self.var_keyword.get(),
            "download_dir": self.var_download.get(),
            "headless": self.var_headless.get(),
            "username": self.var_username.get(),
            "password": self.var_password.get(),
            "role": self.var_role.get(),
            "process_mode": self.var_process_mode.get(),
            "collect_all_pages": self.var_collect_all_pages.get(),
            "max_pages": self.var_max_pages.get(),
            "strict_filter": self.var_strict_filter.get(),
            "match_mode": self.var_match_mode.get(),
        }

    def apply_profile_data(self, data: dict) -> None:
        self.var_keyword.set(data.get("keyword", self.var_keyword.get()))
        self.var_download.set(data.get("download_dir", self.var_download.get()))
        self.var_headless.set(True)
        self.var_username.set(data.get("username", self.var_username.get()))
        self.var_password.set(data.get("password", self.var_password.get()))
        self.var_role.set(data.get("role", self.var_role.get()))
        self.var_process_mode.set(data.get("process_mode", self.var_process_mode.get()))
        self.var_collect_all_pages.set(
            bool(data.get("collect_all_pages", self.var_collect_all_pages.get()))
        )
        self.var_max_pages.set(str(data.get("max_pages", self.var_max_pages.get())))
        self.var_strict_filter.set(bool(data.get("strict_filter", self.var_strict_filter.get())))
        self.var_match_mode.set(str(data.get("match_mode", self.var_match_mode.get())))

    def shutdown(self):
        if self.driver is not None:
            try:
                self.driver.quit()
            except Exception:
                pass


class DocumentationFrame(ttk.Frame):
    def __init__(self, master):
        super().__init__(master)
        self.var_template = tk.StringVar(value="")
        self.var_output_dir = tk.StringVar(value=str(Path.cwd() / "generated_docs"))
        self.var_output_name = tk.StringVar(value="tender-document.docx")
        self.var_username = tk.StringVar(value="")
        self.var_role = tk.StringVar(value="procurement_officer")
        self.var_process_mode = tk.StringVar(value="esjn")
        self.audit_file = Path.cwd() / "compliance" / "audit" / "events.jsonl"
        self._build_ui()

    def _build_ui(self):
        top = ttk.Frame(self)
        top.pack(fill="x", padx=8, pady=8)

        ttk.Label(top, text="Template (.docx):").grid(row=0, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.var_template, width=70).grid(row=0, column=1, sticky="we")
        ttk.Button(top, text="Browse...", command=self.choose_template).grid(row=0, column=2, sticky="w", padx=6)

        ttk.Label(top, text="Output folder:").grid(row=1, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.var_output_dir, width=70).grid(row=1, column=1, sticky="we")
        ttk.Button(top, text="Browse...", command=self.choose_output_dir).grid(row=1, column=2, sticky="w", padx=6)

        ttk.Label(top, text="Output file name:").grid(row=2, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.var_output_name, width=40).grid(row=2, column=1, sticky="w")
        ttk.Label(top, text="Username:").grid(row=3, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.var_username, width=30).grid(row=3, column=1, sticky="w")
        ttk.Label(top, text="Role:").grid(row=3, column=2, sticky="w")
        role_box = ttk.Combobox(
            top,
            textvariable=self.var_role,
            values=("procurement_officer", "admin", "compliance_auditor", "viewer"),
            width=24,
            state="readonly",
        )
        role_box.grid(row=3, column=3, sticky="w")
        ttk.Label(top, text="Process mode:").grid(row=3, column=4, sticky="w")
        mode_box = ttk.Combobox(
            top,
            textvariable=self.var_process_mode,
            values=("esjn", "epazar"),
            width=12,
            state="readonly",
        )
        mode_box.grid(row=3, column=5, sticky="w")

        actions = ttk.Frame(self)
        actions.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Button(actions, text="Scan placeholders", command=self.scan_placeholders).pack(side="left")
        ttk.Button(actions, text="Generate document", command=self.generate_document).pack(side="left", padx=6)

        ttk.Label(self, text="Values (KEY=value, one per line):").pack(anchor="w", padx=8, pady=(0, 4))
        self.values_text = tk.Text(self, height=18, wrap="none")
        self.values_text.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        self.values_text.configure(
            bg=MATRIX_BG,
            fg=MATRIX_FG,
            insertbackground=MATRIX_ACCENT,
            selectbackground=MATRIX_ACCENT,
            selectforeground=MATRIX_BG,
            highlightthickness=1,
            highlightbackground=MATRIX_ACCENT,
            relief="flat",
        )

        self.log_text = tk.Text(self, height=8, wrap="word")
        self.log_text.pack(fill="x", padx=8, pady=(0, 8))
        self.log_text.configure(
            bg=MATRIX_BG,
            fg=MATRIX_FG,
            insertbackground=MATRIX_ACCENT,
            selectbackground=MATRIX_ACCENT,
            selectforeground=MATRIX_BG,
            highlightthickness=1,
            highlightbackground=MATRIX_ACCENT,
            relief="flat",
        )
        self.log_text.config(state="disabled")

    def log(self, msg: str):
        self.log_text.config(state="normal")
        self.log_text.insert("end", time.strftime("%H:%M:%S  ") + msg + "\n")
        self.log_text.see("end")
        self.log_text.config(state="disabled")

    def _enforce_runtime_policy(self, action: str) -> bool:
        decision = load_runtime_policy_gate(COMPLIANCE_RULES_DIR).decide(action)
        active = ",".join(decision.active_rule_ids) if decision.active_rule_ids else "none"
        self.log(
            f"POLICY_GATE action={action} module={decision.module} "
            f"allowed={str(decision.allowed).lower()} active_rules={active} reason={decision.reason}"
        )
        if not decision.allowed:
            messagebox.showerror("Compliance gate", decision.reason)
            return False
        return True

    def choose_template(self):
        path = filedialog.askopenfilename(
            title="Choose template",
            filetypes=[("Word template", "*.docx")],
        )
        if path:
            self.var_template.set(path)

    def choose_output_dir(self):
        path = filedialog.askdirectory(initialdir=self.var_output_dir.get() or str(Path.cwd()))
        if path:
            self.var_output_dir.set(path)

    def scan_placeholders(self):
        template_path = self.var_template.get().strip()
        if not template_path:
            messagebox.showwarning("Template missing", "Select a .docx template.")
            return
        try:
            placeholders = extract_placeholders_from_docx(template_path)
            self.values_text.delete("1.0", "end")
            for key in placeholders:
                self.values_text.insert("end", f"{key}=\n")
            self.log(f"FOUND placeholders: {len(placeholders)}")
        except Exception as exc:
            self.log(f"ERROR: {exc}")
            messagebox.showerror("Error", str(exc))

    def _parse_values(self) -> dict[str, str]:
        values: dict[str, str] = {}
        raw = self.values_text.get("1.0", "end").strip()
        for line in raw.splitlines():
            if not line.strip():
                continue
            if "=" not in line:
                raise ValueError(f"Invalid line: {line}")
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip()
        return values

    def generate_document(self):
        template_path = self.var_template.get().strip()
        output_dir = self.var_output_dir.get().strip() or str(Path.cwd() / "generated_docs")
        output_name = self.var_output_name.get().strip() or "tender-document.docx"
        username = (self.var_username.get() or "").strip()
        role = (self.var_role.get() or "").strip()
        if not self._enforce_runtime_policy("generate_document"):
            return
        route = route_action(self.var_process_mode.get(), "generate_document")
        self.log(
            f"WORKFLOW_ROUTE mode={route.mode} action=generate_document "
            f"allowed={str(route.allowed).lower()} message={route.message}"
        )
        if not route.allowed:
            messagebox.showerror("Workflow routing", route.message)
            return
        decision = authorize_action("generate_document", username, role)
        self.log(build_auth_audit_event("generate_document", username, role, decision.allowed, decision.reason))
        if not decision.allowed:
            messagebox.showerror("Authorization denied", decision.reason)
            return
        if not template_path:
            messagebox.showwarning("Template missing", "Select a .docx template.")
            return
        try:
            values = self._parse_values()
            if route.mode == "epazar":
                missing = validate_required_inputs(
                    {
                        "EPAZAR_OPERATOR_ID": values.get("EPAZAR_OPERATOR_ID", ""),
                        "EPAZAR_CATALOG_ITEM": values.get("EPAZAR_CATALOG_ITEM", ""),
                        "EPAZAR_PROCUREMENT_REF": values.get("EPAZAR_PROCUREMENT_REF", ""),
                    },
                    ["EPAZAR_OPERATOR_ID", "EPAZAR_CATALOG_ITEM", "EPAZAR_PROCUREMENT_REF"],
                )
                if missing:
                    raise ValueError(
                        "Missing mandatory ePazar input fields: " + ", ".join(missing)
                    )
            output_path = str(Path(output_dir) / output_name)
            render_docx_template(template_path, output_path, values)
            dossier_ref = (
                values.get("DOSSIER_ID")
                or values.get("TENDER_ID")
                or Path(output_name).stem
            )
            required_attachment_items = [
                (k, v.strip()) for k, v in values.items() if k.startswith("ATTACHMENT_REQUIRED_")
            ]
            missing_attachment_fields = [k for k, v in required_attachment_items if not v]
            if missing_attachment_fields:
                raise ValueError(
                    "Missing required attachment values for: " + ", ".join(sorted(missing_attachment_fields))
                )
            required_attachments = [v for _, v in required_attachment_items]
            pack = create_workspace_pack(
                base_output_dir=output_dir,
                dossier_ref=dossier_ref,
                primary_document_path=output_path,
                required_attachment_paths=required_attachments,
            )
            self.log(f"GENERATED: {output_path}")
            self.log(f"WORKSPACE: {pack['workspace_dir']}")
            self.log(f"CHECKLIST: {pack['checklist_path']}")
            append_audit_event(
                audit_file=self.audit_file,
                event_type="generate_document",
                actor=username or "anonymous",
                module="doc_builder",
                status="success",
                dossier_id=str(dossier_ref),
                metadata={
                    "template_path": template_path,
                    "output_path": output_path,
                    "workspace_dir": pack["workspace_dir"],
                    "required_attachment_count": len(required_attachments),
                },
            )
            messagebox.showinfo("Done", f"Document saved:\n{output_path}")
        except ValueError as exc:
            self.log(f"VALIDATION_ERROR: {exc}")
            guidance = build_corrective_guidance(
                error_code="validation_missing_fields",
                action="generate_document",
                mode=self.var_process_mode.get(),
            )
            self.log(
                f"GUIDANCE code={guidance['error_code']} retry_safe={str(guidance['retry_safe']).lower()}"
            )
            for step in guidance["steps"]:
                self.log(f"GUIDANCE_STEP: {step}")
            append_audit_event(
                audit_file=self.audit_file,
                event_type="generate_document",
                actor=username or "anonymous",
                module="doc_builder",
                status="failed",
                dossier_id=None,
                metadata={
                    "template_path": template_path,
                    "output_dir": output_dir,
                    "output_name": output_name,
                    "error_message": str(exc),
                },
            )
            messagebox.showerror("Validation error", str(exc))
        except Exception as exc:
            self.log(f"ERROR: {exc}")
            guidance = build_corrective_guidance(
                error_code="unexpected_error",
                action="generate_document",
                mode=self.var_process_mode.get(),
            )
            self.log(
                f"GUIDANCE code={guidance['error_code']} retry_safe={str(guidance['retry_safe']).lower()}"
            )
            for step in guidance["steps"]:
                self.log(f"GUIDANCE_STEP: {step}")
            append_audit_event(
                audit_file=self.audit_file,
                event_type="generate_document",
                actor=username or "anonymous",
                module="doc_builder",
                status="failed",
                dossier_id=None,
                metadata={
                    "template_path": template_path,
                    "output_dir": output_dir,
                    "output_name": output_name,
                    "error_message": str(exc),
                },
            )
            messagebox.showerror("Error", str(exc))

    def get_profile_data(self) -> dict:
        return {
            "template": self.var_template.get(),
            "output_dir": self.var_output_dir.get(),
            "output_name": self.var_output_name.get(),
            "username": self.var_username.get(),
            "role": self.var_role.get(),
            "process_mode": self.var_process_mode.get(),
            "values_text": self.values_text.get("1.0", "end"),
        }

    def apply_profile_data(self, data: dict) -> None:
        self.var_template.set(data.get("template", self.var_template.get()))
        self.var_output_dir.set(data.get("output_dir", self.var_output_dir.get()))
        self.var_output_name.set(data.get("output_name", self.var_output_name.get()))
        self.var_username.set(data.get("username", self.var_username.get()))
        self.var_role.set(data.get("role", self.var_role.get()))
        self.var_process_mode.set(data.get("process_mode", self.var_process_mode.get()))
        self.values_text.delete("1.0", "end")
        self.values_text.insert("1.0", data.get("values_text", ""))


class ProcurementsApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Public Procurements App")
        self.geometry("1200x780")
        self.minsize(1000, 680)
        self.var_include_password = tk.BooleanVar(value=False)
        self.var_matrix_rain = tk.BooleanVar(value=True)
        self._apply_matrix_theme()
        self._rain_job = None
        self._rain_drops: list[dict] = []
        self._rain_chars = "01ABCDEFGHJKLMNPQRSTUVWXYZ"

        self.header_canvas = tk.Canvas(
            self,
            height=84,
            bg=MATRIX_BG,
            highlightthickness=0,
            bd=0,
        )
        self.header_canvas.pack(fill="x", padx=8, pady=(8, 0))
        self.after(80, self._init_matrix_rain)

        top = ttk.Frame(self)
        top.pack(fill="x", padx=8, pady=(8, 0))
        ttk.Button(top, text="Load profile", command=self.load_profile).pack(side="left")
        ttk.Button(top, text="Save profile", command=self.save_profile).pack(side="left", padx=6)
        ttk.Checkbutton(
            top,
            text="Include password in saved profile",
            variable=self.var_include_password,
        ).pack(side="left", padx=(12, 0))
        ttk.Checkbutton(
            top,
            text="Matrix rain",
            variable=self.var_matrix_rain,
            command=self._toggle_matrix_rain,
        ).pack(side="left", padx=(12, 0))

        notebook = ttk.Notebook(self)
        notebook.pack(fill="both", expand=True)

        self.search_tab = TenderSearchFrame(notebook)
        self.docs_tab = DocumentationFrame(notebook)
        notebook.add(self.search_tab, text="Tender Search")
        notebook.add(self.docs_tab, text="Documentation Builder")

        self.protocol("WM_DELETE_WINDOW", self._on_close)

    def _init_matrix_rain(self):
        if not self.var_matrix_rain.get():
            return
        width = max(400, self.header_canvas.winfo_width())
        cols = max(18, width // 24)
        self._rain_drops = []
        for i in range(cols):
            self._rain_drops.append(
                {
                    "x": 12 + i * 24,
                    "y": random.randint(-120, 40),
                    "speed": random.randint(4, 10),
                    "len": random.randint(6, 14),
                }
            )
        self._animate_matrix_rain()

    def _animate_matrix_rain(self):
        if not self.var_matrix_rain.get():
            self.header_canvas.delete("rain")
            self._rain_job = None
            return
        self.header_canvas.delete("rain")
        height = self.header_canvas.winfo_height()
        if height < 30:
            height = 84

        for drop in self._rain_drops:
            head_char = random.choice(self._rain_chars)
            for i in range(drop["len"]):
                y = drop["y"] - i * 12
                if y < -20 or y > height + 20:
                    continue
                if i == 0:
                    color = "#D7FFD7"
                elif i < 3:
                    color = MATRIX_ACCENT
                else:
                    color = "#2A8F3A"
                ch = head_char if i == 0 else random.choice(self._rain_chars)
                self.header_canvas.create_text(
                    drop["x"],
                    y,
                    text=ch,
                    fill=color,
                    font=("Consolas", 11, "bold" if i == 0 else "normal"),
                    tags=("rain",),
                )
            drop["y"] += drop["speed"]
            if drop["y"] - drop["len"] * 12 > height + 40:
                drop["y"] = random.randint(-120, -20)
                drop["speed"] = random.randint(4, 10)
                drop["len"] = random.randint(6, 14)

        self._rain_job = self.after(80, self._animate_matrix_rain)

    def _toggle_matrix_rain(self):
        if self.var_matrix_rain.get():
            if self._rain_job is None:
                self._init_matrix_rain()
        else:
            if self._rain_job is not None:
                try:
                    self.after_cancel(self._rain_job)
                except Exception:
                    pass
                self._rain_job = None
            self.header_canvas.delete("rain")

    def _apply_matrix_theme(self):
        self.configure(bg=MATRIX_BG)
        self.option_add("*Font", "Consolas 10")
        self.option_add("*TCombobox*Listbox.background", MATRIX_BG)
        self.option_add("*TCombobox*Listbox.foreground", MATRIX_FG)

        style = ttk.Style(self)
        style.theme_use("clam")

        style.configure(".", background=MATRIX_PANEL, foreground=MATRIX_FG, fieldbackground=MATRIX_BG)
        style.configure("TFrame", background=MATRIX_PANEL)
        style.configure("TLabel", background=MATRIX_PANEL, foreground=MATRIX_FG)
        style.configure(
            "TButton",
            background=MATRIX_PANEL_ALT,
            foreground=MATRIX_ACCENT,
            borderwidth=1,
            focusthickness=1,
            focuscolor=MATRIX_ACCENT,
            padding=6,
        )
        style.map(
            "TButton",
            background=[("active", "#0E250E"), ("pressed", "#123012")],
            foreground=[("active", "#B8FFB8"), ("pressed", "#D7FFD7")],
        )
        style.configure(
            "TCheckbutton",
            background=MATRIX_PANEL,
            foreground=MATRIX_FG,
            indicatorcolor=MATRIX_BG,
        )
        style.map("TCheckbutton", foreground=[("active", MATRIX_ACCENT)])

        style.configure(
            "TEntry",
            fieldbackground=MATRIX_BG,
            foreground=MATRIX_FG,
            insertcolor=MATRIX_ACCENT,
            bordercolor=MATRIX_ACCENT,
            lightcolor=MATRIX_ACCENT,
            darkcolor=MATRIX_ACCENT,
            padding=4,
        )

        style.configure(
            "TNotebook",
            background=MATRIX_BG,
            borderwidth=0,
            tabmargins=[2, 2, 2, 0],
        )
        style.configure(
            "TNotebook.Tab",
            background=MATRIX_PANEL_ALT,
            foreground=MATRIX_MUTED,
            padding=[12, 6],
            borderwidth=1,
        )
        style.map(
            "TNotebook.Tab",
            background=[("selected", "#123012"), ("active", "#0E250E")],
            foreground=[("selected", MATRIX_ACCENT), ("active", "#B8FFB8")],
        )

        style.configure(
            "Treeview",
            background=MATRIX_BG,
            fieldbackground=MATRIX_BG,
            foreground=MATRIX_FG,
            bordercolor=MATRIX_ACCENT,
            rowheight=24,
        )
        style.map(
            "Treeview",
            background=[("selected", "#103010")],
            foreground=[("selected", "#D7FFD7")],
        )
        style.configure(
            "Treeview.Heading",
            background=MATRIX_PANEL_ALT,
            foreground=MATRIX_ACCENT,
            relief="flat",
            borderwidth=1,
        )
        style.map(
            "Treeview.Heading",
            background=[("active", "#143214")],
            foreground=[("active", "#D7FFD7")],
        )

        style.configure(
            "Vertical.TScrollbar",
            troughcolor=MATRIX_BG,
            background=MATRIX_PANEL_ALT,
            arrowcolor=MATRIX_ACCENT,
            bordercolor=MATRIX_ACCENT,
        )

    def _collect_profile(self) -> dict:
        search_data = self.search_tab.get_profile_data()
        if not self.var_include_password.get():
            search_data["password"] = ""
        return {
            "version": 1,
            "search": search_data,
            "docs": self.docs_tab.get_profile_data(),
            "meta": {
                "password_included": bool(self.var_include_password.get()),
            },
        }

    def _apply_profile(self, data: dict) -> None:
        self.search_tab.apply_profile_data(data.get("search", {}))
        self.docs_tab.apply_profile_data(data.get("docs", {}))

    def save_profile(self):
        path = filedialog.asksaveasfilename(
            title="Save profile",
            defaultextension=".json",
            filetypes=[("JSON", "*.json")],
            initialfile="procurements-profile.json",
        )
        if not path:
            return
        if self.var_include_password.get():
            proceed = messagebox.askyesno(
                "Confirm password export",
                "You are about to save the password in plain text inside the JSON profile. Continue?",
            )
            if not proceed:
                return
        try:
            profile = self._collect_profile()
            Path(path).write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
            messagebox.showinfo("Saved", f"Profile saved:\n{path}")
        except Exception as exc:
            messagebox.showerror("Error", str(exc))

    def load_profile(self):
        path = filedialog.askopenfilename(
            title="Load profile",
            filetypes=[("JSON", "*.json")],
        )
        if not path:
            return
        try:
            data = json.loads(Path(path).read_text(encoding="utf-8"))
            self._apply_profile(data)
            messagebox.showinfo("Loaded", f"Profile loaded:\n{path}")
        except Exception as exc:
            messagebox.showerror("Error", str(exc))

    def _on_close(self):
        if self._rain_job is not None:
            try:
                self.after_cancel(self._rain_job)
            except Exception:
                pass
        self.search_tab.shutdown()
        self.destroy()


if __name__ == "__main__":
    ProcurementsApp().mainloop()
